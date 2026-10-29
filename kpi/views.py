import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Sum, Avg, Count, Q
from django.core.exceptions import PermissionDenied
from decimal import Decimal
from django.db import transaction
from datetime import date
import calendar
from .models import PerformancePeriod as EvaluationCycle
from .models import (
    PerformancePeriod, PerformanceReview, MonthlyKPI, 
    DailyWork, TaskEvidence, WorkApproval, KPI, KPIAssignment, RatingScale, AssignedTask
)
from .forms import (
    KPIForm, KPIAssignmentForm, PerformancePeriodForm,
    DailyWorkForm, SupervisorApprovalForm, AssignedTaskForm
)
from employees.models import Employee, Department, Designation, DailyFloorAttendance, MonthlyMigration, FloorDisciplinaryAction, FloorProductionKPI
from employees.views import get_current_employee, get_current_employee, get_user_role


def get_user_role(user):
    if user.is_superuser:
        return 'SUPER_ADMIN'
    if hasattr(user, 'role_profile'):
        return user.role_profile.role
    return 'EMPLOYEE'


def ensure_periods_exist_for_work_dates():
    """
    Scans all daily work entries and automatically creates 
    any missing PerformancePeriod cycles (e.g., July 2026, August 2026, etc.)
    """
    work_dates = DailyWork.objects.values_list('work_date', flat=True).distinct()
    for w_date in work_dates:
        if w_date:
            year = w_date.year
            month = w_date.month
            month_name = calendar.month_name[month]
            period_name = f"{month_name} {year}"
            
            _, last_day = calendar.monthrange(year, month)
            start_d = date(year, month, 1)
            end_d = date(year, month, last_day)

            PerformancePeriod.objects.get_or_create(
                name=period_name,
                defaults={
                    'period_type': 'Monthly',
                    'start_date': start_d,
                    'end_date': end_d,
                    'status': 'Open',
                    'is_locked': False
                }
            )


# ==========================================
# 1. KPI Configuration & Period Setup
# ==========================================
@login_required
def kpi_configuration(request):
    """
    Manager / Supervisor / Admin Hub to configure KPI Templates,
    assign KPI weights to staff, and create evaluation periods.
    """
    role = get_user_role(request.user)
    current_employee = getattr(request.user, 'employee_profile', None)

    active_tab = request.GET.get('tab', 'templates')
    kpis = KPI.objects.select_related('department', 'designation').all()
    periods = PerformancePeriod.objects.all()
    
    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        assignments = KPIAssignment.objects.select_related('employee', 'kpi', 'period').all()
        assignable_employees = Employee.objects.filter(is_active=True)
    elif role == 'DEPT_HEAD' and current_employee:
        assignments = KPIAssignment.objects.filter(employee__department=current_employee.department)
        assignable_employees = Employee.objects.filter(department=current_employee.department, is_active=True)
    elif role == 'SUPERVISOR' and current_employee:
        assignments = KPIAssignment.objects.filter(employee__supervisor=current_employee)
        assignable_employees = Employee.objects.filter(supervisor=current_employee, is_active=True)
    else:
        assignments = KPIAssignment.objects.filter(employee=current_employee) if current_employee else KPIAssignment.objects.none()
        assignable_employees = Employee.objects.none()

    if request.method == 'POST':
        action = request.POST.get('action')

        if action == 'add_kpi':
            form = KPIForm(request.POST)
            if form.is_valid():
                form.save()
                messages.success(request, "KPI Template created successfully.")
            else:
                messages.error(request, "Please check the KPI form values.")
            return redirect(f"{reverse('kpi_list')}?tab=templates")

        elif action == 'assign_kpi':
            emp_id = request.POST.get('employee')
            kpi_id = request.POST.get('kpi')
            period_id = request.POST.get('period')
            weight = Decimal(request.POST.get('weight', 0))
            target = Decimal(request.POST.get('target', 0))
            unit = request.POST.get('unit', 'Count')
            effective_date = request.POST.get('effective_date') or timezone.now().date()

            if emp_id and kpi_id and period_id:
                emp = Employee.objects.get(id=emp_id)
                kpi_obj = KPI.objects.get(id=kpi_id)
                period_obj = PerformancePeriod.objects.get(id=period_id)

                current_total = KPIAssignment.objects.filter(
                    employee=emp, period=period_obj
                ).exclude(kpi=kpi_obj).aggregate(total=Sum('weight'))['total'] or Decimal('0.0')

                if (current_total + weight) > Decimal('100.0'):
                    messages.error(request, f"Cannot assign {weight}%. Total employee KPI weight would reach {current_total + weight}% (Max 100%).")
                else:
                    KPIAssignment.objects.update_or_create(
                        employee=emp, kpi=kpi_obj, period=period_obj,
                        defaults={
                            'weight': weight,
                            'target': target,
                            'unit': unit,
                            'effective_date': effective_date,
                            'assigned_by': request.user,
                        }
                    )
                    messages.success(request, f"KPI '{kpi_obj.name}' assigned to {emp.full_name} ({weight}%).")
            return redirect(f"{reverse('kpi_list')}?tab=assignments")

        elif action == 'add_period':
            p_form = PerformancePeriodForm(request.POST)
            if p_form.is_valid():
                p_form.save()
                messages.success(request, "Evaluation Period registered successfully.")
            else:
                messages.error(request, "Error creating performance period.")
            return redirect(f"{reverse('kpi_list')}?tab=periods")

    context = {
        'kpis': kpis,
        'assignments': assignments,
        'periods': periods,
        'assignable_employees': assignable_employees,
        'kpi_form': KPIForm(),
        'period_form': PerformancePeriodForm(initial={'start_date': timezone.now().date(), 'end_date': timezone.now().date()}),
        'active_tab': active_tab,
        'can_configure': role in ['SUPER_ADMIN', 'HR_ADMIN', 'DEPT_HEAD', 'SUPERVISOR'],
    }
    return render(request, 'kpi/kpi_list.html', context)


# ==========================================
# 2. Daily Work Operations
# ==========================================
@login_required
def daily_work_list(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    if not current_employee and role not in ['SUPER_ADMIN', 'HR_ADMIN']:
        return redirect('dashboard')

    # Employee views their own work entries
    if current_employee:
        works = DailyWork.objects.filter(employee=current_employee)
    else:
        works = DailyWork.objects.all()

    # Query Parameters
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()
    status_filter = request.GET.get('status', '').strip()

    if start_date:
        works = works.filter(work_date__gte=start_date)
    if end_date:
        works = works.filter(work_date__lte=end_date)
    if status_filter:
        works = works.filter(status=status_filter)

    works = works.select_related('kpi', 'approval').order_by('-work_date', '-id')

    # Summary metrics for the filtered range
    total_entries = works.count()
    total_verified = works.filter(status='Approved').count()
    total_pending = works.filter(status='Submitted').count()

    return render(request, 'kpi/daily_work_list.html', {
        'works': works,
        'start_date': start_date,
        'end_date': end_date,
        'status_filter': status_filter,
        'total_entries': total_entries,
        'total_verified': total_verified,
        'total_pending': total_pending,
        'title': 'Daily Work Entries'
    })


@login_required
def daily_work_create(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    if not current_employee and role not in ['SUPER_ADMIN', 'HR_ADMIN']:
        messages.error(request, "No employee profile linked to your user account.")
        return redirect('dashboard')

    task_id = request.GET.get('task_id')
    kpi_id = request.GET.get('kpi_id')
    task_title = request.GET.get('title')
    target_val = request.GET.get('target')

    assigned_task = None
    if task_id:
        assigned_task = AssignedTask.objects.filter(id=task_id).first()

    if request.method == 'POST':
        form = DailyWorkForm(request.POST, request.FILES, employee=current_employee)
        if form.is_valid():
            with transaction.atomic():
                daily_work = form.save(commit=False)
                daily_work.employee = current_employee
                daily_work.status = 'Submitted'
                daily_work.submitted_at = timezone.now()

                # If submitting an assigned task, link it and update status
                posted_task_id = request.POST.get('assigned_task_id')
                if posted_task_id:
                    task_obj = AssignedTask.objects.filter(id=posted_task_id).first()
                    if task_obj:
                        if hasattr(daily_work, 'assigned_task'):
                            daily_work.assigned_task = task_obj
                        task_obj.status = 'Submitted'
                        task_obj.save()

                daily_work.save()

                # Handle evidence upload
                if 'evidence_file' in request.FILES:
                    file_obj = request.FILES['evidence_file']
                    TaskEvidence.objects.create(
                        daily_work=daily_work,
                        file=file_obj,
                        original_filename=file_obj.name,
                        uploaded_by=request.user,
                        uploaded_at=timezone.now()
                    )

                # Determine the fallback approver
                approver = None
                if current_employee:
                    approver = (
                        current_employee.supervisor 
                        or current_employee.department_head 
                        or current_employee
                    )

                approval_defaults = {
                    'verified_achievement': daily_work.achievement_quantity,
                    'quality_score': Decimal('100.0'),
                    'timeliness_score': Decimal('100.0'),
                    'supervisor_score': Decimal('100.0'),
                }
                
                work_approval_fields = [f.name for f in WorkApproval._meta.get_fields()]
                if 'approved_by' in work_approval_fields:
                    approval_defaults['approved_by'] = approver
                elif 'approver' in work_approval_fields:
                    approval_defaults['approver'] = approver
                elif 'supervisor' in work_approval_fields:
                    approval_defaults['supervisor'] = approver

                WorkApproval.objects.get_or_create(
                    daily_work=daily_work,
                    defaults=approval_defaults
                )

            messages.success(request, f"Daily work '{daily_work.task_title}' submitted successfully for supervisor approval.")
            return redirect('daily_work_list')
        else:
            for field, errors in form.errors.items():
                for err in errors:
                    messages.error(request, f"{field.replace('_', ' ').title()}: {err}")
    else:
        initial_data = {'work_date': timezone.now().date()}
        if kpi_id:
            initial_data['kpi'] = kpi_id
        if task_title:
            initial_data['task_title'] = task_title
        if target_val:
            initial_data['target_quantity'] = target_val
            initial_data['achievement_quantity'] = target_val
        if assigned_task:
            initial_data['task_title'] = assigned_task.title
            initial_data['target_quantity'] = assigned_task.target
            initial_data['achievement_quantity'] = assigned_task.target
            if assigned_task.kpi_id:
                initial_data['kpi'] = assigned_task.kpi_id

        form = DailyWorkForm(initial=initial_data, employee=current_employee)

    return render(request, 'kpi/daily_work_form.html', {
        'form': form,
        'assigned_task': assigned_task,
        'title': 'New Daily Work Entry'
    })


@login_required
def daily_work_edit(request, work_id):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    work = get_object_or_404(
        DailyWork.objects.select_related('employee', 'kpi', 'approval'), 
        id=work_id
    )

    if role not in ['SUPER_ADMIN', 'HR_ADMIN']:
        if not current_employee or work.employee_id != current_employee.id:
            raise PermissionDenied("You can only edit your own daily work submissions.")

    if work.status == 'Approved' and role not in ['SUPER_ADMIN', 'HR_ADMIN']:
        messages.warning(request, "Approved entries are locked. Contact your supervisor to return it if changes are needed.")
        return redirect('daily_work_list')

    if request.method == 'POST':
        form = DailyWorkForm(request.POST, request.FILES, instance=work, employee=work.employee)
        if form.is_valid():
            with transaction.atomic():
                updated_work = form.save(commit=False)
                updated_work.status = 'Submitted'
                updated_work.submitted_at = timezone.now()
                updated_work.save()

                if 'evidence_file' in request.FILES:
                    file_obj = request.FILES['evidence_file']
                    TaskEvidence.objects.create(
                        daily_work=updated_work,
                        file=file_obj,
                        original_filename=file_obj.name,
                        uploaded_by=request.user,
                        uploaded_at=timezone.now()
                    )

                if hasattr(updated_work, 'approval'):
                    approval = updated_work.approval
                    approval.verified_achievement = updated_work.achievement_quantity
                    approval.save()

            messages.success(request, f"Daily work '{updated_work.task_title}' has been updated and resubmitted for verification.")
            return redirect('daily_work_list')
        else:
            for field, err_list in form.errors.items():
                for err in err_list:
                    messages.error(request, f"{field.replace('_', ' ').title()}: {err}")
    else:
        form = DailyWorkForm(instance=work, employee=work.employee)

    supervisor_feedback = None
    if hasattr(work, 'approval') and work.approval.approval_remarks:
        supervisor_feedback = work.approval.approval_remarks

    return render(request, 'kpi/daily_work_form.html', {
        'form': form,
        'work': work,
        'supervisor_feedback': supervisor_feedback,
        'title': f'Edit & Resubmit Daily Work #{work.id}'
    })


# ==========================================
# 3. Calculation Logic
# ==========================================
def calculate_employee_period_kpi(employee, period):
    """
    Core Mathematical Engine: Aggregates floor attendance, migration logs,
    floor production KPI rankings, and approved daily tasks into a final monthly scorecard.
    """
    assignments = KPIAssignment.objects.filter(employee=employee, period=period)
    total_system_weighted_score = Decimal('0.0')
    emp_floor = getattr(employee, 'floor', None)

    for assign in assignments:
        kpi_name_lower = assign.kpi.name.lower()
        kpi_code_lower = assign.kpi.code.lower()

        # CASE 1: Floor-based Absenteeism KPI
        if ('absent' in kpi_name_lower or 'absenteeism' in kpi_code_lower) and emp_floor:
            floor_logs = DailyFloorAttendance.objects.filter(
                floor=emp_floor,
                date__range=(period.start_date, period.end_date)
            )
            if floor_logs.exists():
                avg_absent_rate = floor_logs.aggregate(Avg('absenteeism_rate'))['absenteeism_rate__avg'] or Decimal('0.0')
                target_threshold = Decimal(str(assign.target or 3.00))
                total_verified = round(Decimal(str(avg_absent_rate)), 2)
                effective_target = target_threshold

                if avg_absent_rate <= target_threshold:
                    achievement_pct = Decimal('100.0')
                else:
                    achievement_pct = max(Decimal('0.0'), round((target_threshold / Decimal(str(avg_absent_rate))) * Decimal('100.0'), 2))
            else:
                total_verified = Decimal('0.0')
                effective_target = assign.target
                achievement_pct = Decimal('0.0')

            avg_quality = Decimal('100.0')
            avg_timeliness = Decimal('100.0')
            avg_sup = Decimal('100.0')

        # CASE 2: Migration / Turnover KPI
        elif ('migration' in kpi_name_lower or 'turnover' in kpi_code_lower) and emp_floor:
            migration_records = MonthlyMigration.objects.filter(
                floor=emp_floor,
                month__range=(period.start_date, period.end_date)
            )
            target_limit = Decimal(str(assign.target or 5.00))
            if migration_records.exists():
                avg_migration_rate = migration_records.aggregate(Avg('migration_rate'))['migration_rate__avg'] or Decimal('0.0')
                total_verified = round(Decimal(str(avg_migration_rate)), 2)
                effective_target = target_limit

                if avg_migration_rate <= target_limit:
                    achievement_pct = Decimal('100.0')
                else:
                    achievement_pct = max(Decimal('0.0'), round((target_limit / Decimal(str(avg_migration_rate))) * Decimal('100.0'), 2))
            else:
                total_verified = Decimal('0.0')
                effective_target = target_limit
                achievement_pct = Decimal('0.0')

            avg_quality = Decimal('100.0')
            avg_timeliness = Decimal('100.0')
            avg_sup = Decimal('100.0')

        # CASE 3: Floor Manpower Optimization
        elif ('optimization of manpower' in kpi_name_lower or 'manpower attendance' in kpi_name_lower) and emp_floor:
            floor_logs = DailyFloorAttendance.objects.filter(
                floor=emp_floor,
                date__range=(period.start_date, period.end_date)
            )
            if floor_logs.exists():
                tot_present = floor_logs.aggregate(Sum('present_count'))['present_count__sum'] or 0
                tot_manpower = floor_logs.aggregate(Sum('total_manpower'))['total_manpower__sum'] or 0
                total_verified = Decimal(str(tot_present))
                effective_target = Decimal(str(tot_manpower))
                
                if tot_manpower > 0:
                    achievement_pct = min(Decimal('100.0'), round((Decimal(tot_present) / Decimal(tot_manpower)) * Decimal('100.0'), 2))
                else:
                    achievement_pct = Decimal('0.0')
            else:
                total_verified = Decimal('0.0')
                effective_target = assign.target
                achievement_pct = Decimal('0.0')

            avg_quality = Decimal('100.0')
            avg_timeliness = Decimal('100.0')
            avg_sup = Decimal('100.0')

        # CASE 4: Standard Individual Daily Work Submission KPIs
        else:
            approved_entries = DailyWork.objects.filter(
                employee=employee,
                kpi=assign.kpi,
                work_date__range=(period.start_date, period.end_date),
                status__in=['Approved', 'Partially Approved']
            )

            total_target = approved_entries.aggregate(Sum('target_quantity'))['target_quantity__sum'] or Decimal('0.0')
            total_verified = Decimal('0.0')
            avg_quality = Decimal('0.0')
            avg_timeliness = Decimal('0.0')
            avg_sup = Decimal('0.0')

            count = approved_entries.count()
            if count > 0:
                approvals = WorkApproval.objects.filter(daily_work__in=approved_entries)
                total_verified = approvals.aggregate(Sum('verified_achievement'))['verified_achievement__sum'] or Decimal('0.0')
                avg_quality = approvals.aggregate(Avg('quality_score'))['quality_score__avg'] or Decimal('0.0')
                avg_timeliness = approvals.aggregate(Avg('timeliness_score'))['timeliness_score__avg'] or Decimal('0.0')
                avg_sup = approvals.aggregate(Avg('supervisor_score'))['supervisor_score__avg'] or Decimal('0.0')

            effective_target = total_target if total_target > Decimal('0.0') else assign.target
            if effective_target > Decimal('0.0'):
                achievement_pct = min(Decimal('100.0'), (Decimal(str(total_verified)) / Decimal(str(effective_target))) * Decimal('100.0'))
            else:
                achievement_pct = Decimal('100.0') if total_verified > Decimal('0.0') else Decimal('0.0')

        weighted_score = round((achievement_pct * Decimal(str(assign.weight))) / Decimal('100.0'), 2)
        total_system_weighted_score += weighted_score

        MonthlyKPI.objects.update_or_create(
            employee=employee,
            period=period,
            kpi=assign.kpi,
            defaults={
                'weight': assign.weight,
                'total_target': round(effective_target, 2),
                'total_verified_achievement': round(total_verified, 2),
                'achievement_percentage': round(achievement_pct, 2),
                'average_quality_score': round(avg_quality, 2),
                'average_timeliness_score': round(avg_timeliness, 2),
                'average_supervisor_score': round(avg_sup, 2),
                'weighted_score': round(weighted_score, 2),
                'final_kpi_score': round(weighted_score, 2),
            }
        )

    # -------------------------------------------------------------
    # CASE 5: Check Floor Production KPI (In-Charge / Production Roles)
    # -------------------------------------------------------------
    floor_prod_logs = FloorProductionKPI.objects.filter(
        Q(employee=employee) | (Q(floor=emp_floor) if emp_floor else Q()),
        date__range=(period.start_date, period.end_date)
    )

    if floor_prod_logs.exists():
        prod_agg = floor_prod_logs.aggregate(
            avg_score=Avg('kpi_rank_score'),
            tot_target=Sum('target_quantity'),
            tot_produced=Sum('produced_quantity'),
            avg_eff=Avg('efficiency_rate'),
            avg_dhu=Avg('dhu_rate'),
        )
        avg_prod_score = Decimal(str(round(prod_agg['avg_score'] or 0, 2)))
        tot_target = prod_agg['tot_target'] or 0
        tot_produced = prod_agg['tot_produced'] or 0
        achieve_rate = round((Decimal(tot_produced) / Decimal(tot_target) * 100), 2) if tot_target > 0 else Decimal('0.00')

        # Ensure a designated KPI template exists in the system
        floor_kpi, _ = KPI.objects.get_or_create(
            code='PROD_FLOOR_OPS',
            defaults={
                'name': 'Floor Production & Operational Quality Score',
                'description': 'Composite score evaluated from Daily Floor Production, Efficiency, DHU, and Shipments.',
                'measurement_type': 'Percentage',
                'default_weight': Decimal('100.00'),
                'default_target': Decimal('100.00'),
                'unit': 'Percent'
            }
        )

        # Weight: 100% if no other KPIs assigned, otherwise shares 50%
        prod_weight = Decimal('100.00') if assignments.count() == 0 else Decimal('50.00')
        weighted_prod_score = round((avg_prod_score * prod_weight) / Decimal('100.00'), 2)

        MonthlyKPI.objects.update_or_create(
            employee=employee,
            period=period,
            kpi=floor_kpi,
            defaults={
                'weight': prod_weight,
                'total_target': Decimal(str(tot_target)),
                'total_verified_achievement': Decimal(str(tot_produced)),
                'achievement_percentage': achieve_rate,
                'average_quality_score': Decimal(str(round(prod_agg['avg_dhu'] or 0, 2))),
                'average_timeliness_score': Decimal(str(round(prod_agg['avg_eff'] or 0, 2))),
                'average_supervisor_score': Decimal('100.00'),
                'weighted_score': weighted_prod_score,
                'final_kpi_score': avg_prod_score,
            }
        )

        if assignments.count() == 0:
            total_system_weighted_score = avg_prod_score
        else:
            total_system_weighted_score = round((total_system_weighted_score / 2) + weighted_prod_score, 2)

    # Save to PerformanceReview
    review, _ = PerformanceReview.objects.get_or_create(employee=employee, period=period)
    
    if not review.is_overridden:
        review.system_generated_score = round(total_system_weighted_score, 2)
        if review.supervisor_evaluation_score == Decimal('0.0'):
            review.supervisor_evaluation_score = round(total_system_weighted_score, 2)
            
        final = (review.system_generated_score * Decimal('0.80')) + (review.supervisor_evaluation_score * Decimal('0.20'))
        review.final_score = round(final, 2)

        matching_rating = RatingScale.objects.filter(
            min_score__lte=review.final_score, 
            max_score__gte=review.final_score
        ).first()
        review.rating = matching_rating
        review.save()
        
    return review


# ==========================================
# 4. Supervisor Approvals
# ==========================================
@login_required
def approval_inbox(request):
    role = get_user_role(request.user)
    current_employee = getattr(request.user, 'employee_profile', None)

    if role not in ['SUPER_ADMIN', 'HR_ADMIN', 'DEPT_HEAD', 'SUPERVISOR']:
        raise PermissionDenied("General employees cannot access the approval inbox.")

    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        pending_items = DailyWork.objects.filter(status='Submitted').select_related('employee', 'kpi')
    elif role == 'DEPT_HEAD' and current_employee:
        pending_items = DailyWork.objects.filter(employee__department=current_employee.department, status='Submitted').select_related('employee', 'kpi')
    elif role == 'SUPERVISOR' and current_employee:
        pending_items = DailyWork.objects.filter(employee__supervisor=current_employee, status='Submitted').select_related('employee', 'kpi')
    else:
        pending_items = DailyWork.objects.none()

    return render(request, 'kpi/approval_inbox.html', {'pending_items': pending_items})


@login_required
def approval_process(request, work_id):
    work = get_object_or_404(DailyWork.objects.select_related('employee', 'kpi'), id=work_id)
    current_employee = getattr(request.user, 'employee_profile', None)

    if request.method == 'POST':
        form = SupervisorApprovalForm(request.POST)
        if form.is_valid():
            approval, _ = WorkApproval.objects.get_or_create(
                daily_work=work,
                defaults={'approved_by': current_employee, 'verified_achievement': work.achievement_quantity}
            )
            
            approval.verified_achievement = form.cleaned_data['verified_achievement']
            approval.quality_score = form.cleaned_data['quality_score']
            approval.timeliness_score = form.cleaned_data['timeliness_score']
            approval.supervisor_score = form.cleaned_data['supervisor_score']
            approval.approval_remarks = form.cleaned_data['approval_remarks']
            approval.approved_by = current_employee
            approval.approved_at = timezone.now()
            approval.save()

            work.status = form.cleaned_data['approval_status']
            work.save()

            matching_period = PerformancePeriod.objects.filter(
                start_date__lte=work.work_date,
                end_date__gte=work.work_date
            ).first()
            
            if matching_period:
                calculate_employee_period_kpi(work.employee, matching_period)

            messages.success(request, f"Work item #{work.id} verified as '{work.status}' and KPI recalculated.")
            return redirect('approval_inbox')
    else:
        initial_data = {
            'verified_achievement': work.achievement_quantity,
            'quality_score': 100,
            'timeliness_score': 100,
            'supervisor_score': 100,
        }
        form = SupervisorApprovalForm(initial=initial_data)

    return render(request, 'kpi/approval_detail.html', {'work': work, 'form': form})


# ==========================================
# 5. Multi-Period Automated KPI Aggregation
# ==========================================
@login_required
def monthly_kpi_summary(request):
    ensure_periods_exist_for_work_dates()

    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    # Handle Modal "Add Month" POST request if triggered from UI
    if request.method == 'POST' and request.POST.get('action') == 'create_period':
        if user.is_superuser:
            p_year = int(request.POST.get('period_year', 2026))
            p_month = int(request.POST.get('period_month', 9))
            start_d = date(p_year, p_month, 1)
            last_day = calendar.monthrange(p_year, p_month)[1]
            end_d = date(p_year, p_month, last_day)

            period_title = f"{start_d.strftime('%B %Y')} • Monthly (1 Month) ({start_d.strftime('%d-%b-%Y')} to {end_d.strftime('%d-%b-%Y')})"
            new_p, created = PerformancePeriod.objects.get_or_create(
                start_date=start_d,
                end_date=end_d,
                defaults={'name': period_title, 'period_type': 'Monthly', 'status': 'Open'}[cite: 2]
            )
            messages.success(request, f"Period '{new_p.name}' ready.")
            return redirect(f"{request.path}?period={new_p.id}")

    periods = PerformancePeriod.objects.all().order_by('-start_date')
    period_id = request.GET.get('period')
    target_emp_id = request.GET.get('employee_id')

    # Scope selectable employees by role
    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        selectable_employees = Employee.objects.filter(is_active=True).select_related('department', 'designation')
    elif role == 'DEPT_HEAD' and current_employee:
        selectable_employees = Employee.objects.filter(
            department=current_employee.department, is_active=True
        ).select_related('department', 'designation')
    elif role == 'SUPERVISOR' and current_employee:
        selectable_employees = Employee.objects.filter(
            Q(supervisor=current_employee) | Q(id=current_employee.id), is_active=True
        ).select_related('department', 'designation')
    elif current_employee:
        selectable_employees = Employee.objects.filter(id=current_employee.id)
    else:
        selectable_employees = Employee.objects.none()

    # Determine active target employee
    if target_emp_id:
        target_employee = selectable_employees.filter(id=target_emp_id).first()
    elif current_employee:
        target_employee = current_employee
    else:
        target_employee = selectable_employees.first()

    # Determine selected period
    if period_id:
        selected_period = get_object_or_404(PerformancePeriod, id=period_id)
    elif target_employee:
        latest_work = DailyWork.objects.filter(employee=target_employee).order_by('-work_date').first()
        if latest_work:
            selected_period = PerformancePeriod.objects.filter(
                start_date__lte=latest_work.work_date,
                end_date__gte=latest_work.work_date
            ).first() or periods.first()
        else:
            selected_period = periods.first()
    else:
        selected_period = periods.first()

    review = None
    kpi_breakdown = []

    # Calculate and Aggregate
    if target_employee and selected_period:
        # Run unified engine (handles DailyWork + Attendance + Floor Production)
        review = calculate_employee_period_kpi(target_employee, selected_period)

        # Retrieve distinct MonthlyKPI breakdown
        kpi_breakdown = list(
            MonthlyKPI.objects.filter(
                employee=target_employee, 
                period=selected_period
            ).select_related('kpi').order_by('kpi__code')
        )

    return render(request, 'kpi/monthly_kpi_summary.html', {
        'periods': periods,
        'selected_period': selected_period,
        'selectable_employees': selectable_employees,
        'target_employee': target_employee,
        'review': review,
        'kpi_breakdown': kpi_breakdown,
        'user_role': role,
    })


# ==========================================
# 6. Reports, Preview & Export
# ==========================================
@login_required
def consolidated_performance_report(request):
    ensure_periods_exist_for_work_dates()

    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    periods = PerformancePeriod.objects.all().order_by('-start_date')
    selected_period_id = request.GET.get('period')
    selected_dept_id = request.GET.get('department')
    selected_grade = request.GET.get('grade')
    search_q = request.GET.get('q', '').strip()

    if selected_period_id:
        selected_period = get_object_or_404(PerformancePeriod, id=selected_period_id)
    else:
        selected_period = periods.first()

    rating_scales = RatingScale.objects.all()

    if selected_period:
        assigned_emp_ids = (
            KPIAssignment.objects.filter(period=selected_period)
            .values_list('employee_id', flat=True)
            .distinct()
        )
        employees_to_calc = Employee.objects.filter(
            id__in=assigned_emp_ids, is_active=True
        )
        if not employees_to_calc.exists():
            employees_to_calc = Employee.objects.filter(is_active=True)

        for emp in employees_to_calc:
            calculate_employee_period_kpi(emp, selected_period)

    base_reviews = PerformanceReview.objects.filter(
        period=selected_period
    ).select_related(
        'employee__department',
        'employee__designation',
        'employee__supervisor',
        'rating',
    )

    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        reviews_qs = base_reviews
        departments = Department.objects.filter(is_active=True)
    elif role == 'DEPT_HEAD' and current_employee:
        reviews_qs = base_reviews.filter(
            employee__department=current_employee.department
        )
        departments = Department.objects.filter(id=current_employee.department_id)
    elif role == 'SUPERVISOR' and current_employee:
        reviews_qs = base_reviews.filter(
            Q(employee__supervisor=current_employee) | Q(employee=current_employee)
        )
        departments = Department.objects.filter(id=current_employee.department_id)
    elif current_employee:
        reviews_qs = base_reviews.filter(employee=current_employee)
        departments = Department.objects.filter(id=current_employee.department_id)
    else:
        reviews_qs = base_reviews.none()
        departments = Department.objects.none()

    if selected_dept_id and role in ['SUPER_ADMIN', 'HR_ADMIN']:
        reviews_qs = reviews_qs.filter(employee__department_id=selected_dept_id)
    if selected_grade:
        reviews_qs = reviews_qs.filter(rating__grade=selected_grade)
    if search_q:
        reviews_qs = reviews_qs.filter(
            Q(employee__employee_code__icontains=search_q)
            | Q(employee__first_name__icontains=search_q)
            | Q(employee__last_name__icontains=search_q)
        )

    dept_performance = []
    if role in ['SUPER_ADMIN', 'HR_ADMIN', 'DEPT_HEAD']:
        for dept in departments:
            dept_reviews = base_reviews.filter(employee__department=dept)
            emp_count = dept_reviews.count()
            avg_score = dept_reviews.aggregate(avg=Avg('final_score'))['avg'] or Decimal('0.0')
            dept_performance.append({
                'department': dept,
                'count': emp_count,
                'avg_score': round(avg_score, 2),
            })

    total_evaluated = reviews_qs.count()
    overall_avg = reviews_qs.aggregate(avg=Avg('final_score'))['avg'] or Decimal('0.0')
    outstanding_count = reviews_qs.filter(final_score__gte=Decimal('95.0')).count()
    needs_improvement_count = reviews_qs.filter(final_score__lt=Decimal('70.0')).count()

    context = {
        'periods': periods,
        'selected_period': selected_period,
        'departments': departments,
        'selected_dept_id': selected_dept_id,
        'rating_scales': rating_scales,
        'selected_grade': selected_grade,
        'search_q': search_q,
        'reviews': reviews_qs,
        'dept_performance': dept_performance,
        'total_evaluated': total_evaluated,
        'overall_avg': round(overall_avg, 2),
        'outstanding_count': outstanding_count,
        'needs_improvement_count': needs_improvement_count,
        'user_role': role,
        'can_edit_kpi': role in ['SUPER_ADMIN', 'HR_ADMIN', 'DEPT_HEAD', 'SUPERVISOR'],
    }
    return render(request, 'kpi/consolidated_report.html', context)


@login_required
def edit_performance_review(request, review_id):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    review = get_object_or_404(
        PerformanceReview.objects.select_related(
            'employee__department', 'employee__designation', 'employee__supervisor', 'period', 'rating'
        ),
        id=review_id
    )

    if role not in ['SUPER_ADMIN', 'HR_ADMIN']:
        if role == 'DEPT_HEAD' and current_employee:
            if review.employee.department_id != current_employee.department_id:
                raise PermissionDenied("You can only edit KPIs for your own department.")
        elif role == 'SUPERVISOR' and current_employee:
            if review.employee.supervisor_id != current_employee.id:
                raise PermissionDenied("You can only edit KPIs for your direct subordinates.")
        else:
            raise PermissionDenied("Employees cannot edit KPI reviews.")

    if request.method == 'POST':
        sup_score = Decimal(request.POST.get('supervisor_evaluation_score', review.supervisor_evaluation_score))
        sup_comments = request.POST.get('supervisor_comments', '').strip()
        hr_comments = request.POST.get('hr_comments', '').strip()
        
        is_override = request.POST.get('is_overridden') == 'on'
        override_score = request.POST.get('override_score')
        override_reason = request.POST.get('override_reason', '').strip()

        review.supervisor_evaluation_score = min(Decimal('100.0'), max(Decimal('0.0'), sup_score))
        review.supervisor_comments = sup_comments
        review.hr_comments = hr_comments

        if is_override and role in ['SUPER_ADMIN', 'HR_ADMIN'] and override_score:
            review.is_overridden = True
            if not review.original_score:
                review.original_score = review.final_score
            review.final_score = Decimal(override_score)
            review.override_reason = override_reason
            review.overridden_by = user
        else:
            review.final_score = (review.system_generated_score * Decimal('0.80')) + (review.supervisor_evaluation_score * Decimal('0.20'))

        matching_rating = RatingScale.objects.filter(
            min_score__lte=review.final_score, max_score__gte=review.final_score
        ).first()
        review.rating = matching_rating
        review.save()

        messages.success(request, f"KPI Performance Review for {review.employee.full_name} updated successfully.")
        return redirect('performance_reviews')

    return render(request, 'kpi/review_edit.html', {
        'review': review,
        'user_role': role,
    })


@login_required
def export_consolidated_kpi_excel(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    period_id = request.GET.get('period')
    dept_id = request.GET.get('department')

    period = get_object_or_404(PerformancePeriod, id=period_id) if period_id else PerformancePeriod.objects.first()
    reviews = PerformanceReview.objects.filter(period=period).select_related(
        'employee__department', 'employee__designation', 'employee__supervisor', 'rating'
    )

    if role == 'DEPT_HEAD' and current_employee:
        reviews = reviews.filter(employee__department=current_employee.department)
    elif role == 'SUPERVISOR' and current_employee:
        reviews = reviews.filter(Q(employee__supervisor=current_employee) | Q(employee=current_employee))
    elif role == 'EMPLOYEE' and current_employee:
        reviews = reviews.filter(employee=current_employee)
    elif dept_id and role in ['SUPER_ADMIN', 'HR_ADMIN']:
        reviews = reviews.filter(employee__department_id=dept_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "KPI Consolidated Report"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    meta_font = Font(name="Calibri", size=11, bold=True, color="475569")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.merge_cells("A1:K1")
    ws["A1"] = "KPIMaster - Consolidated KPI Performance Scorecard"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Evaluation Period: {period.name} ({period.get_period_type_display()})"
    ws["A2"].font = meta_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    headers = [
        "SL", "Employee Code", "Employee Name", "Department", "Designation", 
        "Supervisor", "System Score (80%)", "Supervisor Score (20%)", 
        "Final Score (%)", "Grade", "Rating Status"
    ]

    ws.append([])
    ws.append(headers)
    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, rev in enumerate(reviews, 1):
        row_data = [
            idx,
            rev.employee.employee_code,
            rev.employee.full_name,
            rev.employee.department.name,
            rev.employee.designation.title,
            rev.employee.supervisor.full_name if rev.employee.supervisor else "—",
            float(rev.system_generated_score),
            float(rev.supervisor_evaluation_score),
            float(rev.final_score),
            rev.rating.grade if rev.rating else "N/A",
            rev.rating.name if rev.rating else "Not Graded"
        ]
        ws.append(row_data)

        row_idx = 4 + idx
        for col_idx in range(1, 12):
            c = ws.cell(row=row_idx, column=col_idx)
            c.font = data_font
            c.border = border
            if col_idx in [1, 2, 7, 8, 9, 10]:
                c.alignment = Alignment(horizontal="center")
            if col_idx == 9:
                c.font = bold_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="KPI_Report_{period.name.replace(" ", "_")}.xlsx"'
    wb.save(response)
    return response


@login_required
def consolidated_report_preview(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    period_id = request.GET.get('period')
    periods = PerformancePeriod.objects.all()
    selected_period = get_object_or_404(PerformancePeriod, id=period_id) if period_id else (periods.filter(status='Open').first() or periods.first())

    reviews_qs = PerformanceReview.objects.filter(period=selected_period).select_related(
        'employee__department', 'employee__designation', 'employee__supervisor', 'rating'
    )

    if role == 'DEPT_HEAD' and current_employee:
        reviews_qs = reviews_qs.filter(employee__department=current_employee.department)
    elif role == 'SUPERVISOR' and current_employee:
        reviews_qs = reviews_qs.filter(Q(employee__supervisor=current_employee) | Q(employee=current_employee))
    elif role == 'EMPLOYEE' and current_employee:
        reviews_qs = reviews_qs.filter(employee=current_employee)

    total_evaluated = reviews_qs.count()
    overall_avg = reviews_qs.aggregate(avg=Avg('final_score'))['avg'] or Decimal('0.0')

    context = {
        'selected_period': selected_period,
        'reviews': reviews_qs,
        'total_evaluated': total_evaluated,
        'overall_avg': round(overall_avg, 2),
        'generated_at': timezone.now(),
        'generated_by': request.user,
    }
    return render(request, 'kpi/report_preview_pdf.html', context)


@login_required
def single_employee_kpi_card_preview(request, review_id):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    review = get_object_or_404(
        PerformanceReview.objects.select_related(
            'employee__department', 'employee__designation', 'employee__supervisor', 'period', 'rating'
        ),
        id=review_id
    )

    if role == 'EMPLOYEE' and current_employee and review.employee_id != current_employee.id:
        raise PermissionDenied("You can only view your own KPI card.")
    elif role == 'SUPERVISOR' and current_employee and (review.employee.supervisor_id != current_employee.id and review.employee_id != current_employee.id):
        raise PermissionDenied
    elif role == 'DEPT_HEAD' and current_employee and review.employee.department_id != current_employee.department_id:
        raise PermissionDenied

    kpi_breakdown = MonthlyKPI.objects.filter(
        employee=review.employee,
        period=review.period
    ).select_related('kpi')

    return render(request, 'kpi/employee_card_preview_pdf.html', {
        'review': review,
        'kpi_breakdown': kpi_breakdown,
        'generated_at': timezone.now(),
        'generated_by': request.user,
    })


# ==========================================
# 7. Task & Team Work Views
# ==========================================
@login_required
def assign_task_view(request):
    if request.method == 'POST':
        form = AssignedTaskForm(request.POST)
        if form.is_valid():
            task = form.save(commit=False)
            task.assigned_by = request.user
            task.status = 'Assigned'
            task.save()

            messages.success(request, f"Task '{task.title}' successfully assigned to {task.employee.full_name}.")
            return redirect('my_tasks')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f"{field.replace('_', ' ').title()}: {error}")
    else:
        form = AssignedTaskForm()

    return render(request, 'kpi/assign_task.html', {
        'form': form,
        'title': 'Assign Task to Employee'
    })


@login_required
def my_tasks(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        tasks = AssignedTask.objects.all().select_related('employee', 'kpi', 'assigned_by').order_by('-assigned_date')
    elif current_employee:
        tasks = AssignedTask.objects.filter(employee=current_employee).select_related('kpi', 'assigned_by').order_by('-assigned_date')
    else:
        tasks = AssignedTask.objects.none()

    return render(request, 'kpi/my_tasks.html', {
        'tasks': tasks,
        'title': 'My Assigned Tasks'
    })


@login_required
def team_work_list(request):
    user = request.user
    role = get_user_role(user)
    current_employee = getattr(user, 'employee_profile', None)

    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        subordinates = Employee.objects.filter(is_active=True)
        works = DailyWork.objects.all()
    elif current_employee:
        subordinates_qs = Employee.objects.filter(supervisor=current_employee, is_active=True)
        if role in ['DEPT_HEAD', 'MANAGER'] and current_employee.department:
            dept_members = Employee.objects.filter(department=current_employee.department, is_active=True)
            subordinates = (subordinates_qs | dept_members).distinct()
        else:
            subordinates = subordinates_qs.distinct()

        if not subordinates.exists():
            subordinates = Employee.objects.filter(id=current_employee.id)

        works = DailyWork.objects.filter(employee__in=subordinates)
    else:
        subordinates = Employee.objects.none()
        works = DailyWork.objects.none()

    selected_emp = request.GET.get('employee', '').strip()
    selected_status = request.GET.get('status', '').strip()
    selected_date = request.GET.get('date', '').strip()

    if selected_emp:
        works = works.filter(employee_id=selected_emp)
    if selected_status:
        works = works.filter(status=selected_status)
    if selected_date:
        works = works.filter(work_date=selected_date)

    works = works.select_related('employee', 'kpi', 'approval').prefetch_related('evidences').order_by('-work_date', '-id')

    return render(request, 'kpi/team_work_list.html', {
        'works': works,
        'subordinates': subordinates,
        'selected_emp': selected_emp,
        'selected_status': selected_status,
        'selected_date': selected_date,
        'title': 'Team Work Submissions'
    })

@login_required
def reports_dashboard(request):
    """
    Central Reports Hub - Overview of all analytical, operational, 
    and performance reports available in KPIMaster.
    """
    today = timezone.now().date()
    current_month_start = today.replace(day=1)

    # Quick system counters
    total_employees = Employee.objects.filter(is_active=True).count()
    attendance_records_month = DailyFloorAttendance.objects.filter(date__gte=current_month_start).count()
    disciplinary_records_month = FloorDisciplinaryAction.objects.filter(date__gte=current_month_start).count()
    production_records_month = FloorProductionKPI.objects.filter(date__gte=current_month_start).count()

    report_categories = [
        {
            'category_title': 'KPI & Performance Evaluations',
            'icon': 'bi-award-fill text-warning',
            'reports': [
                {
                    'title': 'Consolidated KPI Performance Review',
                    'url_name': 'performance_reviews',
                    'icon': 'bi-trophy text-warning',
                    'badge': 'Core KPI',
                    'badge_class': 'bg-warning text-dark',
                    'desc': 'Department-wise composite performance ratings, 80/20 weighted scores, and review ledgers.',
                    'export_types': ['PDF', 'Excel', 'Print']
                },
                {
                    'title': 'Monthly KPI Scores Summary',
                    'url_name': 'monthly_kpi_summary',
                    'icon': 'bi-graph-up-arrow text-primary',
                    'badge': 'Monthly',
                    'badge_class': 'bg-primary text-white',
                    'desc': 'Target vs achievement tracking, individual employee KPI scoring, and achievement percentiles.',
                    'export_types': ['View', 'Print']
                },
                {
                    'title': 'Daily Work & Task Approvals Log',
                    'url_name': 'team_work_list',
                    'icon': 'bi-check2-square text-success',
                    'badge': 'Audit',
                    'badge_class': 'bg-success text-white',
                    'desc': 'Supervisor approval ledger, verified daily claims, and pending task completion status.',
                    'export_types': ['Filter', 'Audit']
                }
            ]
        },
        {
            'category_title': 'Floor Operations & Production Analytics',
            'icon': 'bi-buildings-fill text-info',
            'reports': [
                {
                    'title': 'Floor Attendance & Turnover Analytics',
                    'url_name': 'floor_analytics_summary',
                    'icon': 'bi-bar-chart-steps text-info',
                    'badge': 'Analytics',
                    'badge_class': 'bg-info text-dark',
                    'desc': 'Daily manpower averages, floor-level absenteeism %, joining/resignation turnover rates, and trend charts.',
                    'export_types': ['1-Page PDF', 'Print']
                },
                {
                    'title': 'Floor Disciplinary Actions & Incident Tracker',
                    'url_name': 'floor_disciplinary_analytics',
                    'icon': 'bi-shield-slash text-danger',
                    'badge': 'Compliance',
                    'badge_class': 'bg-danger text-white',
                    'desc': 'Floor incident frequency, severity distribution (warnings, suspensions, dismissals), and incident rates %.',
                    'export_types': ['1-Page PDF', 'Print']
                },
                {
                    'title': 'Floor Production KPI & Assessment Ranking',
                    'url_name': 'floor_production_manage',
                    'icon': 'bi-speedometer2 text-primary',
                    'badge': 'Ranking',
                    'badge_class': 'bg-dark text-white',
                    'desc': '10-Indicator composite leaderboard: Target Achievement, DHU %, Extra OT, Timely Shipment, and Floor Management.',
                    'export_types': ['Real-time', 'Filter']
                }
            ]
        },
        {
            'category_title': 'Workforce & Organization Directory',
            'icon': 'bi-people-fill text-secondary',
            'reports': [
                {
                    'title': 'Employee Master Directory & Placement',
                    'url_name': 'employee_list',
                    'icon': 'bi-people text-primary',
                    'badge': f'{total_employees} Active',
                    'badge_class': 'bg-secondary text-white',
                    'desc': 'Full workforce roster categorized by factory unit, department, line, designation, supervisor, and job grade.',
                    'export_types': ['Directory', 'Profiles']
                },
                {
                    'title': 'Organization Structure & Hierarchy Map',
                    'url_name': 'org_structure',
                    'icon': 'bi-building text-dark',
                    'badge': 'Structure',
                    'badge_class': 'bg-light text-dark border',
                    'desc': 'Multi-unit setup: factory units, departments, sections, sub-sections, lines, and operational floors.',
                    'export_types': ['Hierarchy', 'Config']
                }
            ]
        }
    ]

    return render(request, 'kpi/reports_dashboard.html', {
        'report_categories': report_categories,
        'total_employees': total_employees,
        'attendance_records_month': attendance_records_month,
        'disciplinary_records_month': disciplinary_records_month,
        'production_records_month': production_records_month,
        'title': 'Reports & Analytics Central Hub'
    })


@login_required
def performance_reviews(request):
    """
    Consolidated KPI Performance Review Ledger View with Filter by Period, Department, and Employee.
    """
    ensure_periods_exist_for_work_dates()

    role = get_user_role(request.user)
    current_employee = get_current_employee(request.user)

    periods = PerformancePeriod.objects.all().order_by('-start_date')
    selected_period_id = request.GET.get('period')
    selected_dept = request.GET.get('department', '').strip()
    selected_grade = request.GET.get('grade', '').strip()
    query = request.GET.get('q', '').strip()

    # Default to the most recent period (e.g. September 2026)
    if selected_period_id:
        selected_period = get_object_or_404(PerformancePeriod, id=selected_period_id)
    else:
        selected_period = periods.first()

    # Ensure all active employees for this period are evaluated
    if selected_period:
        active_employees = Employee.objects.filter(is_active=True)
        if selected_dept:
            active_employees = active_employees.filter(department_id=selected_dept)

        # Run scoring calculation so reviews exist in DB
        for emp in active_employees:
            calculate_employee_period_kpi(emp, selected_period)

    # Base query for reviews
    reviews_qs = PerformanceReview.objects.filter(
        period=selected_period
    ).select_related(
        'employee__department',
        'employee__designation',
        'employee__supervisor',
        'rating',
        'period'
    ).order_by('-final_score')

    # Department and Role scoping
    departments = Department.objects.filter(is_active=True)
    if role in ['SUPER_ADMIN', 'HR_ADMIN']:
        pass
    elif role == 'DEPT_HEAD' and current_employee:
        reviews_qs = reviews_qs.filter(employee__department=current_employee.department)
        departments = Department.objects.filter(id=current_employee.department_id)
    elif role == 'SUPERVISOR' and current_employee:
        reviews_qs = reviews_qs.filter(
            Q(employee__supervisor=current_employee) | Q(employee=current_employee)
        )
        departments = Department.objects.filter(id=current_employee.department_id)
    elif current_employee:
        reviews_qs = reviews_qs.filter(employee=current_employee)
        departments = Department.objects.filter(id=current_employee.department_id)
    else:
        reviews_qs = PerformanceReview.objects.none()

    # User Filters
    if selected_dept and role in ['SUPER_ADMIN', 'HR_ADMIN']:
        reviews_qs = reviews_qs.filter(employee__department_id=selected_dept)
    if selected_grade:
        reviews_qs = reviews_qs.filter(rating__grade=selected_grade)
    if query:
        reviews_qs = reviews_qs.filter(
            Q(employee__employee_code__icontains=query) |
            Q(employee__first_name__icontains=query) |
            Q(employee__last_name__icontains=query)
        )

    context = {
        'reviews': reviews_qs,
        'periods': periods,
        'selected_period': selected_period,
        'departments': departments,
        'selected_dept': selected_dept,
        'selected_grade': selected_grade,
        'query': query,
        'user_role': role,
        'title': 'Employee Monthly Performance Reviews'
    }
    return render(request, 'kpi/performance_reviews.html', context)

@login_required
def create_monthly_evaluation_cycle(request):
    if not request.user.is_superuser:
        messages.error(request, "Access restricted to Administrators.")
        return redirect('dashboard')

    if request.method == 'POST':
        year = int(request.POST.get('year'))
        month = int(request.POST.get('month'))
        
        start_date = date(year, month, 1)
        last_day = calendar.monthrange(year, month)[1]
        end_date = date(year, month, last_day)
        
        cycle_name = f"{start_date.strftime('%B %Y')} • Monthly (1 Month) ({start_date.strftime('%d-%b-%Y')} to {end_date.strftime('%d-%b-%Y')})"
        
        cycle, created = EvaluationCycle.objects.get_or_create(
            start_date=start_date,
            end_date=end_date,
            defaults={
                'name': cycle_name,
                'cycle_type': 'MONTHLY',
                'is_active': True,
            }
        )
        
        if created:
            messages.success(request, f"Evaluation Cycle '{cycle_name}' created successfully.")
        else:
            messages.warning(request, f"Evaluation Cycle for {start_date.strftime('%B %Y')} already exists.")
            
        return redirect('monthly_kpi_summary')

    return render(request, 'kpi/create_cycle_modal.html')